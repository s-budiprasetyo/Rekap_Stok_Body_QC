"""Export laporan ke file Excel (.xlsx) dengan format mirip template."""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import engine

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")

FILL_GREY = PatternFill("solid", fgColor="D9D9D9")
FILL_GREEN = PatternFill("solid", fgColor="92D050")
FILL_BLUE = PatternFill("solid", fgColor="BDD7EE")
FILL_PINK = PatternFill("solid", fgColor="FF99CC")

DISPLAY_COLS = [
    ("OP100_PERIKSA", "OP100", "PERIKSA", None),
    ("OP105_PERBAIKAN", "OP105", "PERBAIKAN", None),
    ("OP107_TDKSET_IJP", "OP107", "TDK SET iJP", None),
    ("OP107_TDKSET_CRJP", "OP107", "TDK SET CRJP", None),
    ("OP110_GERINDA", "OP110", "GERINDA", FILL_GREEN),
    ("OP115_CEKULANG", "OP115", "CEK ULANG", FILL_BLUE),
    ("OP120_TAP", "OP120", "TAP", FILL_PINK),
    ("OP166_FITTINGST", "OP166", "FITTING & ST", None),
    ("OP166_KIRIM", "OP166", "KIRIM", None),
]


def build_excel(report_df: pd.DataFrame, tanggal_str: str, sesi: str, kota: str = "Gresik") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stok Body Harian QC"

    n_data = len(DISPLAY_COLS)
    last_col = 3 + n_data + 1  # No, Type, ISI + data cols + GrandTotal

    # ---------- Baris 1: nama perusahaan ----------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, "PT. SURYA PERTIWI NUSANTARA").font = Font(bold=True, size=12)
    ws.cell(1, 1).alignment = CENTER

    # ---------- Baris 2: "Gresik, tanggal" rata kanan, di ATAS kotak judul+legalisasi ----------
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    gc = ws.cell(2, 1, f"{kota}, {tanggal_str}")
    gc.font = Font(size=9)
    gc.alignment = Alignment(horizontal="right", vertical="center")

    # ---------- Kotak gabungan: judul (kiri, 3 baris) + legalisasi (kanan, 3 baris) ----------
    box_row0 = 3
    box_row1 = 5  # 3 baris: DIBUAT/.. | blank ttd | INDIRECT/..

    title_end_col = max(1, last_col - 6)  # sisakan 6 kolom terakhir untuk 3 grup legalisasi (2 kolom/grup)
    legal_start_col = title_end_col + 1
    legal_cols = last_col - title_end_col
    base = legal_cols // 3
    rem = legal_cols % 3
    group_sizes = [base + (1 if i < rem else 0) for i in range(3)]

    ws.merge_cells(start_row=box_row0, start_column=1, end_row=box_row1, end_column=title_end_col)
    tc = ws.cell(box_row0, 1, "STOK BODY HARIAN QC TK")
    tc.font = Font(bold=True, size=16)
    tc.alignment = CENTER

    col = legal_start_col
    for label, size in zip(["DIBUAT", "DIPERIKSA", "DISETUJUI"], group_sizes):
        ws.merge_cells(start_row=box_row0, start_column=col, end_row=box_row0, end_column=col + size - 1)
        c = ws.cell(box_row0, col, label)
        c.font = Font(bold=True, size=10)
        c.alignment = CENTER
        col += size

    col = legal_start_col
    for size in group_sizes:
        ws.merge_cells(start_row=box_row0 + 1, start_column=col, end_row=box_row0 + 1, end_column=col + size - 1)
        col += size

    col = legal_start_col
    for label, size in zip(["INDIRECT", "FOREMAN", "SUPERVISOR"], group_sizes):
        ws.merge_cells(start_row=box_row1, start_column=col, end_row=box_row1, end_column=col + size - 1)
        c = ws.cell(box_row1, col, label)
        c.font = Font(bold=True, size=10)
        c.alignment = CENTER
        col += size

    for r in range(box_row0, box_row1 + 1):
        for c in range(1, last_col + 1):
            ws.cell(r, c).border = BORDER

    # ---------- Baris TANGGAL ----------
    tanggal_row = box_row1 + 1
    ws.merge_cells(start_row=tanggal_row, start_column=1, end_row=tanggal_row, end_column=last_col)
    ws.cell(tanggal_row, 1, f"TANGGAL : {tanggal_str} ( {sesi} )").font = Font(bold=True, size=11)

    header_row1 = tanggal_row + 2
    header_row2 = header_row1 + 1

    ws.merge_cells(start_row=header_row1, start_column=1, end_row=header_row2, end_column=1)
    ws.cell(header_row1, 1, "NO")
    ws.merge_cells(start_row=header_row1, start_column=2, end_row=header_row2, end_column=2)
    ws.cell(header_row1, 2, "TYPE")
    ws.merge_cells(start_row=header_row1, start_column=3, end_row=header_row2, end_column=3)
    ws.cell(header_row1, 3, "ISI PALET/TKRT")

    col = 4
    i = 0
    while i < len(DISPLAY_COLS):
        key, op, label, fill = DISPLAY_COLS[i]
        span = 1
        j = i + 1
        while j < len(DISPLAY_COLS) and DISPLAY_COLS[j][1] == op:
            span += 1
            j += 1
        if span > 1:
            ws.merge_cells(start_row=header_row1, start_column=col, end_row=header_row1, end_column=col + span - 1)
        c = ws.cell(header_row1, col, op)
        if fill:
            for cc in range(col, col + span):
                ws.cell(header_row1, cc).fill = fill
        col += span
        i = j
    ws.merge_cells(start_row=header_row1, start_column=col, end_row=header_row2, end_column=col)
    ws.cell(header_row1, col, "GRAND TOTAL")

    col = 4
    for key, op, label, fill in DISPLAY_COLS:
        c = ws.cell(header_row2, col, label)
        if fill:
            c.fill = fill
        col += 1

    for r in range(header_row1, header_row2 + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True, size=9)
            cell.alignment = CENTER
            cell.border = BORDER

    row = header_row2 + 1
    no_counter = 1
    groups = [g for g in engine.GROUP_ORDER if g in report_df["RowGroup"].unique()]
    col_totals = {k: 0 for k, *_ in DISPLAY_COLS}

    for grp in groups:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        gcell = ws.cell(row, 1, grp)
        gcell.font = Font(bold=True)
        gcell.fill = FILL_GREY
        row += 1

        sub = report_df[report_df["RowGroup"] == grp]
        for _, r in sub.iterrows():
            ws.cell(row, 1, no_counter).alignment = CENTER
            ws.cell(row, 2, str(r["RowLabel"])).alignment = LEFT
            isi_val = str(r["ISI"]).strip()
            ws.cell(row, 3, isi_val if isi_val and isi_val != "nan" else "").alignment = CENTER
            col = 4
            for key, op, label, fill in DISPLAY_COLS:
                val = int(r.get(key, 0) or 0)
                cell = ws.cell(row, col, val if val else None)
                cell.alignment = CENTER
                col_totals[key] += val
                col += 1
            ws.cell(row, col, int(r["GrandTotal"])).font = Font(bold=True)
            ws.cell(row, col).alignment = CENTER

            for c in range(1, last_col + 1):
                ws.cell(row, c).border = BORDER
            no_counter += 1
            row += 1

    totals = engine.compute_totals(report_df)
    footer = [("TOTAL PCS 1", totals["pcs1"], False),
              ("TOTAL AKSESORIS", totals["aksesoris"], False),
              ("TOTAL BODY", totals["body"], True)]
    for label, value, show_cols in footer:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        c1 = ws.cell(row, 1, label)
        c1.font = Font(bold=True)
        c1.fill = FILL_GREY
        col = 4
        if show_cols:
            for key, op, lab, fill in DISPLAY_COLS:
                cc = ws.cell(row, col, col_totals[key])
                cc.font = Font(bold=True)
                cc.fill = FILL_GREY
                cc.alignment = CENTER
                col += 1
        else:
            ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=3 + len(DISPLAY_COLS))
            for c in range(4, 4 + len(DISPLAY_COLS)):
                ws.cell(row, c).fill = FILL_GREY
            col = 4 + len(DISPLAY_COLS)
        cc = ws.cell(row, col, value)
        cc.font = Font(bold=True)
        cc.fill = FILL_GREY
        cc.alignment = CENTER
        for c in range(1, last_col + 1):
            ws.cell(row, c).border = BORDER
        row += 1

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 10
    for c in range(4, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

    _write_raw_data_sheet(wb, report_df)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_raw_data_sheet(wb, report_df: pd.DataFrame):
    """Sheet tersembunyi berisi data mentah report_df apa adanya (bukan untuk dilihat
    manusia), supaya bisa dibaca ulang secara programatik oleh fitur 'Ambil Data'
    tanpa perlu parsing ulang layout tabel yang sudah diformat."""
    ws2 = wb.create_sheet("RawData")
    cols = ["RowLabel", "RowGroup", "ISI", "IsAksesoris", "GrandTotal"] + engine.REPORT_COLUMN_KEYS
    ws2.append(cols)
    for _, r in report_df.iterrows():
        row = [r.get("RowLabel"), r.get("RowGroup"), r.get("ISI"), bool(r.get("IsAksesoris")),
               int(r.get("GrandTotal", 0) or 0)]
        row += [int(r.get(k, 0) or 0) for k in engine.REPORT_COLUMN_KEYS]
        ws2.append(row)
    ws2.sheet_state = "hidden"


def read_report_df(xlsx_bytes: bytes) -> pd.DataFrame:
    """Baca ulang report_df dari sheet 'RawData' tersembunyi di file hasil build_excel().
    Dipakai oleh halaman 'Ambil Data' untuk me-render ulang gambar PNG dari Excel yang
    tersimpan otomatis di Google Drive."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    if "RawData" not in wb.sheetnames:
        raise ValueError("File Excel ini tidak punya sheet RawData (bukan hasil generate aplikasi ini).")
    ws = wb["RawData"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Sheet RawData kosong.")
    header = list(rows[0])
    data_rows = [r for r in rows[1:] if r and r[0] is not None]
    df = pd.DataFrame(data_rows, columns=header)
    df["IsAksesoris"] = df["IsAksesoris"].astype(bool)
    for c in engine.REPORT_COLUMN_KEYS + ["GrandTotal"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).astype(int)
    return df
